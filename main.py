import io
import fitz
import time
import pandas as pd
import streamlit as st
import matplotlib.pyplot as plt

from app.components.header import show_header
from app.components.sidebar import show_sidebar
from app.components.cards import show_dashboard_cards
from app.services.pdf_generator import create_pdf
from app.services.parser import extract_text_from_pdf

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from io import BytesIO
from concurrent.futures import ThreadPoolExecutor
from app.services.llm_service import (explain_medical_report,compare_medical_reports)

st.set_page_config(
    page_title="AI Medical Report Explainer",
    page_icon="🏥",
    layout="wide"
)

show_header()
page = show_sidebar()
st.write(page)

if page == "📄 AI Medical Report":
    uploaded_file = st.file_uploader(
        "Upload your Medical Report",
        type=["pdf", "png", "jpg", "jpeg"]
    )
    st.write("Uploader Loaded")
    
    text = ""
       
    if uploaded_file is not None:

        st.success("✅ File Uploaded Successfully!")
        st.write("Filename:", uploaded_file.name)

        st.write("### File Details")
        st.write("📄 File Name:", uploaded_file.name)
        st.write("📦 File Size:", round(uploaded_file.size / 1024, 2), "KB")

        if uploaded_file.type.startswith("image"):
            st.image(
                uploaded_file,
                caption="Uploaded Medical Report",
                use_container_width=True
            )

        elif uploaded_file.type == "application/pdf":

            pdf = fitz.open(
                stream=uploaded_file.read(),
                filetype="pdf"
            )

            for page in pdf:
                text += page.get_text()

            st.subheader("📄 Extracted Text")

            st.text_area(
                "PDF Content",
                text,
                height=300
            )

    if text:

        if "result" not in st.session_state:
            st.session_state["result"]=None
            st.session_state["pdf_file"]=None

        if st.button("🤖 Explain Report with AI"):
            progress_bar = st.progress(0)

            status = st.empty()
            status.info("📄 Reading Medical Report...")
            progress_bar.progress(20)

            time.sleep(0.5)

            status.info("🧠 AI is Analyzing Report...")
            progress_bar.progress(50)

            time.sleep(0.5)

            status.info("📊 Calculating Health Score...")
            progress_bar.progress(80)

            time.sleep(0.5)

            status.info("✅ Preparing Final Report...")
            progress_bar.progress(100)

            result = explain_medical_report(text)

            progress_bar.empty()
            status.empty()
            st.session_state["result"] = result

            rag_sources = result.get("_rag_sources", [])
            if rag_sources:
                st.caption(
                    "🔎 RAG used local medical knowledge from: "
                    + ", ".join(src["source"] for src in rag_sources)
                )

                # st.write(result["abnormal_values"])
                
            summary=f"""AI MEDICAL REPORT SUMMARY

    Health Score: {result['health_score']}/100
    Risk Level: {result['risk_level']}
    Abnormal Tests: {result['abnormal_tests']}
    Normal Tests: {result['normal_tests']}

    Overall Health Summary
    {result['overall_summary']}
    """
            for item in result["abnormal_values"]:
                summary+=f"""\nTest Name: {item['test_name']}\nPatient Value: {item['patient_value']}\nNormal Range: {item['normal_range']}\nExplanation: {item['explanation']}\nPossible Reason: {item['possible_reason']}\nFood Recommendation: {item['food']}\nLifestyle Suggestion: {item['lifestyle']}\nDoctor Consultation: {item['doctor_consultation']}\n"""
            pdf_file=create_pdf(summary)
            st.session_state["result"]=result
            st.session_state["pdf_file"]=pdf_file

        if st.session_state["result"] is not None:
            result=st.session_state["result"]
            pdf_file=st.session_state["pdf_file"]

    # Dashboard
            show_dashboard_cards(
                result["health_score"],
                result["risk_level"],
                result["abnormal_tests"],
                result["normal_tests"]
            )
            st.subheader("📊 Report Overview")
            col1, col2, col3 = st.columns([1,2,1])
            with col2:
                fig, ax = plt.subplots(figsize=(3,3))
            
            labels = ["Normal", "Abnormal"]
            sizes = [
                result["normal_tests"],
                result["abnormal_tests"]
            ]
            
            if sum(sizes) > 0:

                ax.pie(
                    sizes,
                    labels=labels,
                    autopct="%1.1f%%",
                    startangle=90,
                    radius=0.8,
                    textprops={"fontsize": 9}
                )

                ax.axis("equal")
                st.pyplot(fig, use_container_width=False)

            else:
                st.warning("⚠️ No report data available to display chart.")
            
            st.subheader("📊 Health Score")
            score = result["health_score"]

            filled = "█" * (score // 5)
            empty = "░" * (20 - (score // 5))
            st.markdown(f"### `{filled}{empty} {score}%`")
            
            st.subheader("🩺 AI Medical Explanation")
            st.header("📋 Overall Health Summary")
            st.write(result["overall_summary"])
            
            # ==============================
            # 🚨 Risk Detection
            # ==============================
            st.markdown("---")
            st.header("🚨 Risk Detection")
            risk = result["risk_level"]

            if risk == "High":
                st.error("🔴 High Risk")

            elif risk == "Moderate":
                st.warning("🟡 Moderate Risk")

            else:
                st.success("🟢 Low Risk")
                
            st.subheader("🩺 Possible Health Issues")
            for item in result["abnormal_values"]:
                st.write(f"• {item['test_name']}")
            
            st.subheader("👨‍⚕️ Recommended Doctor")
            doctor = "General Physician"
            tests = " ".join(
                [item["test_name"].lower() for item in result["abnormal_values"]]
            )

            if "heart" in tests or "cholesterol" in tests or "hdl" in tests or "ldl" in tests:
                doctor = "Cardiologist"

            elif "vitamin" in tests:
                doctor = "Nutritionist"

            elif "hemoglobin" in tests or "cbc" in tests:
                doctor = "Hematologist"

            st.info(f"Recommended Specialist: **{doctor}**")
            
            # ==============================
            # Search Medical Test
            # ==============================

            st.header("🩸 Abnormal Values")

            st.markdown(
                "<h4 style='margin-bottom:8px;'>🔍 Search Medical Test</h4>",
                unsafe_allow_html=True
            )

            search = st.text_input(
                "",
                placeholder="Example: Vitamin, HDL, ESR...",
                label_visibility="collapsed"
            )

            filtered_items = []

            if search:
                filtered_items = [

                    item for item in result["abnormal_values"]

                    if search.lower() in item["test_name"].lower()

                ]

                if len(filtered_items) == 0:

                    st.warning("❌ No Medical Test Found")

                else:

                    table_data = []

                    for item in filtered_items:

                        risk = item.get("risk", "").lower()

                        if risk == "high":
                            status = "🔴 High"
                        elif risk == "moderate":
                            status = "🟠 Moderate"
                        elif risk == "low":
                            status = "🟢 Low"
                        else:
                            status = "🔴 Abnormal"

                        table_data.append(
                            {
                                "Test Name": item["test_name"],
                                "Patient Value": item["patient_value"],
                                "Normal Range": item["normal_range"],
                                "Status": status
                            }
                        )

                    df = pd.DataFrame(table_data)
                    csv = df.to_csv(index=False).encode("utf-8")
                    
                    st.dataframe(
                        df,
                        use_container_width=True,
                        hide_index=True
                    )
            else:
                st.info("🔍 Search a medical test to view abnormal values.")
                
            # ===========================
            # Detailed Explanation
            # ===========================
            st.markdown("---")
            st.header("📖 Detailed Explanation")
            for item in result["abnormal_values"]:
                with st.expander(f"🧪 {item['test_name']}"):

                    st.write(f"**Patient Value:** {item['patient_value']}")
                    st.write(f"**Normal Range:** {item['normal_range']}")

                    st.info(item["explanation"])
                    st.warning(f"Possible Reason: {item['possible_reason']}")
                    st.success(f"🍎 Food Recommendation: {item['food']}")
                    st.success(f"🏃 Lifestyle: {item['lifestyle']}")
                    st.error(f"👨‍⚕️ Doctor Consultation: {item['doctor_consultation']}")
            
            # ===========================
            # AI Personalized Health Tips
            # ===========================
            st.markdown("---")
            st.header("💡 AI Personalized Health Tips")

            with st.expander("📖 Show Personalized Health Tips", expanded=False):
                st.markdown("### 🩺 AI Recommendations")
                for i, tip in enumerate(result.get("health_tips", []), start=1):
                    st.markdown(f"**{i}.** ✅ {tip}")

                st.markdown("---")

                st.info(
                    "⚠️ These recommendations are AI-generated based on your uploaded medical report. "
                    "Please consult a qualified healthcare professional before making any medical decisions."
                )

                st.info(
                    "⚠️ This AI-generated report is for informational purposes only and should not be considered a medical diagnosis. Please consult a qualified healthcare professional for medical advice."
                )
                
            # ===========================
            # Download Button
            # ===========================  
            
            # ==========================
            # Create Excel Report
            # ==========================

            wb = Workbook()
            ws = wb.active
            ws.title = "Medical Report"

            headers = [
                "Test Name",
                "Patient Value",
                "Normal Range",
                "Risk",
                "Explanation",
                "Possible Reason",
                "Food",
                "Lifestyle",
                "Doctor Consultation"
            ]

            for col, header in enumerate(headers, start=1):
                cell = ws.cell(row=1, column=col)
                cell.value = header
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill(fill_type="solid", fgColor="1F4E78")
                cell.alignment = Alignment(horizontal="center", vertical="center")

            row = 2

            for item in result.get("abnormal_values", []):

                ws.cell(row=row, column=1).value = item.get("test_name", "")
                ws.cell(row=row, column=2).value = item.get("patient_value", "")
                ws.cell(row=row, column=3).value = item.get("normal_range", "")
                ws.cell(row=row, column=4).value = item.get("risk", "")
                ws.cell(row=row, column=5).value = item.get("explanation", "")
                ws.cell(row=row, column=6).value = item.get("possible_reason", "")
                ws.cell(row=row, column=7).value = item.get("food", "")
                ws.cell(row=row, column=8).value = item.get("lifestyle", "")
                ws.cell(row=row, column=9).value = item.get("doctor_consultation", "")

                row += 1

            for column_cells in ws.columns:

                length = max(len(str(cell.value)) if cell.value else 0 for cell in column_cells)

                ws.column_dimensions[get_column_letter(column_cells[0].column)].width = min(length + 5, 40)

            for row_cells in ws.iter_rows():
                for cell in row_cells:
                    cell.alignment = Alignment(
                        wrap_text=True,
                        vertical="top"
                    )

            excel_file = BytesIO()
            wb.save(excel_file)
            excel_file.seek(0)   
                
            with open(pdf_file, "rb") as file:
                st.download_button(
                    label="📥 Download PDF AI Report",
                    data=file,
                    file_name="AI_Medical_Report.pdf",
                    mime="application/pdf"
                )
                st.download_button(
                    label="📊 Download Excel Report",
                    data=excel_file,
                    file_name="AI_Medical_Report.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                )
                
            # =====================================
            # Compare Reports
            # =====================================

if page == "📊 Compare Reports":
    if "comparison_result" not in st.session_state:
        st.session_state["comparison_result"] = None

    st.subheader("📊 Compare Medical Reports")

    report1 = st.file_uploader(
        "Upload Report 1",
        type=["pdf"],
        key="report1"
    )

    report2 = st.file_uploader(
        "Upload Report 2",
        type=["pdf"],
        key="report2"
    )

    if report1 is not None and report2 is not None:
        st.success("✅ Both reports uploaded successfully!")
        compare_btn = st.button("🔍 Compare Reports")

        if compare_btn:

            with st.spinner("Comparing Reports..."):

                report1_text = extract_text_from_pdf(report1)
                report2_text = extract_text_from_pdf(report2)

                comparison_result = compare_medical_reports(
                    report1_text,
                    report2_text
                )
                st.session_state["comparison_result"] = comparison_result
                result1 = comparison_result["report1"]
                result2 = comparison_result["report2"]

                st.subheader("📊 Comparison Summary")
                col1, col2 = st.columns(2)

                with col1:
                    st.markdown("### 📄 Report 1")
                    st.metric("Health Score", f"{result1['health_score']}/100")
                    st.metric("Risk Level", result1["risk_level"])
                    st.metric("Normal Tests", result1["normal_tests"])
                    st.metric("Abnormal Tests", result1["abnormal_tests"])

                with col2:
                    st.markdown("### 📄 Report 2")
                    st.metric("Health Score", f"{result2['health_score']}/100")
                    st.metric("Risk Level", result2["risk_level"])
                    st.metric("Normal Tests", result2["normal_tests"])
                    st.metric("Abnormal Tests", result2["abnormal_tests"])

                st.divider()
                
                st.subheader("🧪 Test-wise Comparison")
                comparison = comparison_result["comparison"]
                comparison_df = pd.DataFrame(comparison)

                st.dataframe(
                    comparison_df,
                    use_container_width=True,
                    hide_index=True
                )